"""
Corrige el informe 'Taller_Arquitectura_C4_MLOps_ERV_Parte2_UWAS12.docx':
  1. Cambia la fecha de entrega de la portada a 22 de octubre de 2026.
  2. Reemplaza los 4 marcadores [INSERTAR FIGURA ...] por las imágenes de los
     diagramas corregidos, dentro de un marco con borde simple negro.
     Si falta una imagen, deja un marco formal con la referencia a la evidencia.
  3. Fuerza todo el texto a color negro.

Uso:
  pip install python-docx pillow
  Pon las imágenes en la carpeta 'figuras/' con los nombres de FIGURAS y ejecuta:
  python corregir_informe.py
"""
import os, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ENTRADA = "Taller_Arquitectura_C4_MLOps_ERV_Parte2_UWAS12.docx"
SALIDA = "Taller_Arquitectura_C4_MLOps_ERV_Parte2_UWAS12_v2.docx"
CARPETA = "figuras"
GRAYSCALE = False   # True convierte las imágenes a escala de grises
ANCHO_MAX = 6.3     # pulgadas
ALTO_MAX = 7.6      # pulgadas
NEGRO = RGBColor(0, 0, 0)

FECHA_ANTIGUA = "28 de octubre de 2026"
FECHA_NUEVA = "22 de octubre de 2026"

FIGURAS = [
    ("FIGURA 1", "figura1_contexto_corregido.png",
     "Figura 1. Diagrama de Contexto C4 #CORREGIDO: ecosistema WellQ en región UK-South, "
     "ingesta mediante URLs SAS y cumplimiento UK GDPR."),
    ("FIGURA 2", "figura2_contenedores_corregido.png",
     "Figura 2. Diagrama de Contenedores C4 #CORREGIDO: plano de tiempo real (< 200 ms) con "
     "Cloud Run, caché Redis y base primaria OLTP, separado del plano batch con Worker Jobs y réplica OLAP."),
    ("FIGURA 3", "figura3_componentes_corregido.png",
     "Figura 3. Diagrama de Componentes C4 #CORREGIDO: API FastAPI con pool OLTP prioritario "
     "y pool OLAP de solo lectura para analítica y auditoría SOAP."),
    ("FIGURA 4", "figura4_pipeline_corregido.png",
     "Figura 4. Pipeline MLOps y Servicios de Inferencia #CORREGIDO: flujo de 8 etapas con "
     "restricción de aislamiento estricto batch frente a tiempo real y estrategia de fallback."),
]
NOTA_EVIDENCIA = "Evidencia generada en la sesión sincrónica del 22 de octubre de 2026 (carpeta /GPI-Taller Arquitectura/)."


def preparar_imagen(ruta):
    from PIL import Image
    im = Image.open(ruta)
    if GRAYSCALE:
        im = im.convert("L")
        ruta = ruta.rsplit(".", 1)[0] + "_gris.png"
        im.save(ruta)
    w, h = im.size
    escala = min(ANCHO_MAX / w, ALTO_MAX / h)
    return ruta, Inches(w * escala)


def texto_negro(run, tam=None, negrita=None, cursiva=None):
    run.font.color.rgb = NEGRO
    if tam: run.font.size = Pt(tam)
    if negrita is not None: run.bold = negrita
    if cursiva is not None: run.italic = cursiva


def iter_parrafos(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for fila in t.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    yield p
    for s in doc.sections:
        for p in s.footer.paragraphs:
            yield p


def reemplazar_texto(p, viejo, nuevo):
    if viejo not in p.text:
        return False
    completo = p.text.replace(viejo, nuevo)
    for r in p.runs[1:]:
        r.text = ""
    p.runs[0].text = completo
    return True


def siguiente_parrafo(tabla):
    nodo = tabla._tbl.getnext()
    from docx.text.paragraph import Paragraph
    return Paragraph(nodo, tabla._parent) if nodo is not None and nodo.tag.endswith("}p") else None


def main():
    doc = Document(ENTRADA)

    # 1. Fecha de entrega
    cambios = sum(reemplazar_texto(p, FECHA_ANTIGUA, FECHA_NUEVA) for p in iter_parrafos(doc))
    print(f"Fecha actualizada en {cambios} lugar(es).")

    # 2. Figuras
    for tabla in doc.tables:
        if len(tabla.rows) != 1 or len(tabla.columns) != 1:
            continue
        celda = tabla.cell(0, 0)
        texto = celda.text
        figura = next((f for f in FIGURAS if "[INSERTAR " + f[0] in texto), None)
        if not figura:
            continue
        clave, archivo, epigrafe = figura
        ruta = os.path.join(CARPETA, archivo)
        for p in celda.paragraphs[1:]:
            p._element.getparent().remove(p._element)
        p = celda.paragraphs[0]
        for r in p.runs:
            r.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if os.path.exists(ruta):
            ruta_final, ancho = preparar_imagen(ruta)
            p.runs[0].add_picture(ruta_final, width=ancho) if p.runs else p.add_run().add_picture(ruta_final, width=ancho)
            print(f"{clave}: imagen insertada ({archivo}).")
        else:
            r = p.add_run(f"{clave} · {epigrafe.split(': ')[0].split('. ', 1)[1]}")
            texto_negro(r, 11, True)
            p2 = celda.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r2 = p2.add_run(f"Imagen no encontrada: {archivo}")
            texto_negro(r2, 9, cursiva=True)
            print(f"{clave}: falta {ruta}; se deja marco formal.")
        # nota de evidencia dentro del marco
        p3 = celda.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(NOTA_EVIDENCIA)
        texto_negro(r3, 8.5, cursiva=True)
        # epígrafe centrado bajo el marco
        cap = siguiente_parrafo(tabla)
        if cap is not None and cap.text.startswith("Figura"):
            for r in cap.runs[1:]:
                r.text = ""
            cap.runs[0].text = epigrafe
            texto_negro(cap.runs[0], 9, cursiva=True)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 3. Todo el texto en negro
    for p in iter_parrafos(doc):
        for r in p.runs:
            r.font.color.rgb = NEGRO

    doc.save(SALIDA)
    print(f"Listo: {SALIDA}")


if __name__ == "__main__":
    main()
